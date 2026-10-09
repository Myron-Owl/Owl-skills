"""Synthetic regression tests for bounded reads, rich text, header diagnosis and recovery."""
import argparse
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.worksheet.table import Table, TableColumn
from edit_common import EditError, digest, load_plan, read_book, table_header_issues
from edit_simple import execute
from edit_batch import execute_batch
from read_workbook import inspect


def run(root):
    rich = root / 'rich.xlsx'
    w = Workbook()
    w.active['A1'] = CellRichText('技', TextBlock(InlineFont(b=True), '能'))
    w.active['B1'] = '数量'
    w.active.append(['短剑', 1])
    t = Table(displayName='Items', ref='A1:B2', tableColumns=[TableColumn(id=1, name='技能'), TableColumn(id=2, name='数量')])
    w.active.add_table(t)
    w.save(rich)
    w.close()
    h = digest(rich)
    info = inspect(rich, 'Sheet', 'A1:B2')
    assert info['cells'][0]['value'] == '技能' and info['table_header_issue_count'] == 0
    assert digest(rich) == h

    def plan(name, file, before=1, after=2, address='B2', preconditions=None):
        p = root / (name + '.json')
        data = dict(workbook=str(file), operations=[dict(sheet='Sheet', range=address, kind='value', before=[[before]], values=[[after]])])
        if preconditions:
            data['preconditions'] = preconditions
        p.write_text(json.dumps(data), encoding='utf-8')
        return p

    p = plan('rich', rich, preconditions=[dict(sheet='Sheet', range='A1', values=[['技能']])])
    execute(load_plan(p))
    w = read_book(rich)
    assert isinstance(w.active['A1'].value, CellRichText) and w.active['A1'].value[1].font.b
    w.active.tables['Items'].tableColumns[0].name = 'OldName'
    w.save(rich)
    w.close()
    h = digest(rich)
    assert inspect(rich)['table_header_issue_count'] == 1
    try:
        execute(load_plan(plan('bad-header', rich, before=2)))
        raise AssertionError('Expected header rejection')
    except EditError as e:
        assert 'Table header' in str(e) and digest(rich) == h

    for scenario in ('success', 'failure', 'concurrent'):
        files, steps = [], []
        for i in range(2):
            f = root / f'{scenario}-{i}.xlsx'
            w = Workbook()
            w.active['B2'] = 1
            w.save(f)
            w.close()
            files.append(f)
            steps.append(dict(engine='simple', plan=str(plan(f'{scenario}-{i}', f))))
        original = [digest(f) for f in files]
        concurrent_hash = []
        def runner(p):
            if p['workbook'] == str(files[1]) and scenario != 'success':
                if scenario == 'concurrent':
                    w = read_book(files[0])
                    w.active['C3'] = 'external change'
                    w.save(files[0])
                    w.close()
                    concurrent_hash.append(digest(files[0]))
                raise EditError('Injected second-file failure')
            return execute(p)
        try:
            result = execute_batch(dict(steps=steps, recovery_directory=str(root / f'recovery-{scenario}')), {'simple': runner})
            assert scenario == 'success' and result['status'] == 'PASS_BATCH'
            for f in files:
                w = read_book(f)
                assert w.active['B2'].value == 2
                w.close()
        except EditError as e:
            assert scenario != 'success'
            assert digest(files[1]) == original[1]
            if scenario == 'failure':
                assert digest(files[0]) == original[0]
                assert e.context['files'][0]['state'] == 'restored'
            else:
                assert digest(files[0]) == concurrent_hash[0]
                assert e.context['files'][0]['state'] == 'recovery_blocked'
    # Fault injection checks native stage/file attribution without starting Excel.
    import edit_native
    native_plan = load_plan(steps[0]['plan'])
    native_plan['operations'][0]['before'] = [[2]]
    native_plan['dependencies'] = [str(files[1])]
    real_session = edit_native.ExcelSession
    for failed_role in ('dependency', 'target'):
        class FailedOpenSession:
            app = None
            def __enter__(self):
                return self
            def __exit__(self, *args):
                return False
            def open(self, path):
                if failed_role == 'dependency' or str(path) == native_plan['workbook']:
                    raise RuntimeError('Injected native open failure')
                return object()
        edit_native.ExcelSession = FailedOpenSession
        try:
            edit_native.execute(native_plan)
            raise AssertionError('Expected native open failure')
        except EditError as e:
            assert e.context['stage'] == 'open_' + failed_role
            expected_file = native_plan['dependencies'][0] if failed_role == 'dependency' else native_plan['workbook']
            assert e.context['failed_file'] == expected_file and not e.context['original_written']
        finally:
            edit_native.ExcelSession = real_session
    print('PASS: bounded reading, rich-text matching/preservation, table diagnosis, batch success/recovery/concurrent protection')


if __name__ == '__main__':
    with TemporaryDirectory(prefix='excel-skill-regression-') as directory:
        run(Path(directory))
