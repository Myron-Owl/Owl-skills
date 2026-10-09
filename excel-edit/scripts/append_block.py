"""Bounded native format-copy/merge support for a new block below existing contents."""
from copy import copy
from edit_common import bounds, cells, require, style, effective_style_array, content, dimension_signature


def validate(spec):
    require(isinstance(spec, dict) and {'sheet', 'range'} <= set(spec) <=
            {'sheet', 'range', 'copy_formats_from', 'merges', 'row_height'}, 'Invalid append_block fields')
    require(isinstance(spec['sheet'], str), 'append_block needs sheet')
    x1, y1, x2, y2 = bounds(spec['range'])
    require((x2-x1+1)*(y2-y1+1) <= 2000, 'Append block limited to 2,000 cells')
    if 'copy_formats_from' in spec:
        a, b, c, d = bounds(spec['copy_formats_from'])
        require((c-a, d-b) == (x2-x1, y2-y1), 'Format source must have the same dimensions')
    occupied = set()
    require(isinstance(spec.get('merges', []), list), 'merges must be a list')
    for address in spec.get('merges', []):
        a, b, c, d = bounds(address)
        require(x1 <= a <= c <= x2 and y1 <= b <= d <= y2, 'Merge outside append block')
        keys = {(sh, cell) for sh, cell, _, _ in cells(spec['sheet'], address)}
        require(not occupied & keys, 'Overlapping merges')
        occupied |= keys
    if 'row_height' in spec:
        require(type(spec['row_height']) in (int, float) and 0 < spec['row_height'] <= 409,
                'Invalid append row height')


def intersects(a, b):
    from openpyxl.utils.cell import range_boundaries
    x1, y1, x2, y2 = range_boundaries(str(a))
    u1, v1, u2, v2 = range_boundaries(str(b))
    return x1 <= u2 and u1 <= x2 and y1 <= v2 and v1 <= y2


def prepare(wb, p):
    spec = p.get('append_block')
    if spec is None:
        return None
    validate(spec)
    require(sum(len(s._cells) for s in wb) <= 10000, 'Large workbook append needs scoped handling')
    require(spec['sheet'] in wb.sheetnames, 'Missing append sheet')
    ws = wb[spec['sheet']]
    require(not ws.protection.sheet, 'Protected append sheet')
    x1, y1, x2, y2 = bounds(spec['range'])
    last = max((c.row for c in ws._cells.values() if c.value is not None), default=0)
    require(y1 > last, 'Append block must be below existing content; use scoped editing otherwise')
    for address in [spec['range']] + ([spec['copy_formats_from']] if 'copy_formats_from' in spec else []):
        require(not any(intersects(address, m) for m in ws.merged_cells.ranges), 'Source/destination overlaps existing merge')
        require(not any(intersects(address, r) for cf in ws.conditional_formatting for r in cf.sqref.ranges), 'Format-copy region has conditional formatting')
        require(not any(intersects(address, t.ref) for t in ws.tables.values()), 'Table region needs scoped handling')
    require(not any(intersects(spec['range'], r) for dv in ws.data_validations.dataValidation for r in dv.sqref.ranges),
            'Append destination has validation rules')
    require(not any(ws[addr].comment or ws[addr].hyperlink for _, addr, _, _ in cells(ws.title, spec['range'])),
            'Append destination has comments or hyperlinks')
    # Capture old-region appearance once for append edits only, not ordinary input edits.
    old = {(s.title, c.coordinate): (content(c), style(c)) for s in wb for c in s._cells.values()
           if c.__class__.__name__ != 'MergedCell'}
    dimensions = {(s.title, kind, key): dimension_signature(value)
                  for s in wb for kind in ('row_dimensions', 'column_dimensions')
                  for key, value in getattr(s, kind).items()
                  if not (s.title == ws.title and kind == 'row_dimensions' and y1 <= key <= y2)}
    if 'copy_formats_from' in spec:
        a, b, _, _ = bounds(spec['copy_formats_from'])
        source = {(dy, dx): copy(effective_style_array(ws.cell(b+dy, a+dx)))
                  for dy in range(y2-y1+1) for dx in range(x2-x1+1)}
        for _, address, dy, dx in cells(ws.title, spec['range']):
            ws[address]._style = source[dy, dx]
    for address in spec.get('merges', []):
        ws.merge_cells(address)
    if 'row_height' in spec:
        for row in range(y1, y2+1):
            ws.row_dimensions[row].height = spec['row_height']
    expected = {addr: style(ws[addr]) for _, addr, _, _ in cells(ws.title, spec['range'])
                if ws[addr].__class__.__name__ != 'MergedCell'}
    return dict(old=old, expected=expected, merges=str(ws.merged_cells), dimensions=dimensions,
                heights={r: ws.row_dimensions[r].height if r in ws.row_dimensions else None for r in range(y1, y2+1)})


def apply(target, spec):
    if spec is None:
        return
    ws = target.Worksheets(spec['sheet'])
    if 'copy_formats_from' in spec:
        ws.Range(spec['copy_formats_from']).Copy()
        ws.Range(spec['range']).PasteSpecial(Paste=-4122)  # formats only
        target.Application.CutCopyMode = False
    for address in spec.get('merges', []):
        ws.Range(address).Merge()
    if 'row_height' in spec:
        ws.Range(spec['range']).RowHeight = spec['row_height']


def verify(wb, p, expected):
    if expected is None:
        return
    spec = p['append_block']
    ws = wb[spec['sheet']]
    changed = {(sh, addr) for op in p['operations'] for sh, addr, _, _ in cells(op['sheet'], op['range'])}
    appended = {(sh, addr) for sh, addr, _, _ in cells(spec['sheet'], spec['range'])}
    for (sheet, kind, key), signature in expected['dimensions'].items():
        if sheet in {s['name'] for s in p.get('new_sheets', [])}:
            continue  # New dimensions are verified by the new-sheet contract.
        dim = getattr(wb[sheet], kind).get(key)
        require(dim is not None and dimension_signature(dim) == signature,
                f'Unexpected dimension change: {sheet}/{kind}/{key}')
    for (sh, addr), value in expected['old'].items():
        if (sh, addr) not in changed | appended:
            require((content(wb[sh][addr]), style(wb[sh][addr])) == value, f'Unexpected old content/style change: {sh}!{addr}')
    for addr, value in expected['expected'].items():
        if (ws.title, addr) not in changed:
            require(style(ws[addr]) == value and ws[addr].value is None, f'Unexpected append content/style: {addr}')
    require(str(ws.merged_cells) == expected['merges'], 'Append merges differ')
    for row, height in expected['heights'].items():
        found = ws.row_dimensions.get(row)
        require((found.height if found else None) == height, f'Append row height differs: {row}')
