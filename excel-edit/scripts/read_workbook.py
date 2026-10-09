"""Read-only, bounded workbook summaries and cell inspection; never starts Excel."""
import argparse
import json
from pathlib import Path
from openpyxl import load_workbook
from edit_common import read_book, table_header_issues, plan_value, cells, bounds, require, check_unlocked


def inspect(path, sheet=None, address=None, cached=False, for_edit=False):
    require(bool(sheet) == bool(address), "Provide both --sheet and --range")
    if address:
        x1, y1, x2, y2 = bounds(address)
        require((x2-x1+1)*(y2-y1+1) <= 1000, "Read at most 1,000 cells per request")
    availability = None
    if for_edit:
        try:
            check_unlocked(Path(path).resolve())
            availability = {"replaceable_now": True}
        except Exception as error:
            availability = {"replaceable_now": False, "reason": str(error)}
    wb = read_book(path)
    cache = None
    try:
        result = {"workbook": str(Path(path).resolve()), "assurance": "read_only_parse; not native-open or calculation validation",
                  "sheets": [{"name": w.title, "rows": w.max_row, "columns": w.max_column,
                              "tables": [{"name": t.name, "range": t.ref} for t in w.tables.values()]} for w in wb]}
        if for_edit:
            result["edit_availability"] = availability
        issues = table_header_issues(wb)
        result["table_header_issue_count"] = len(issues)
        result["table_header_issues"] = issues[:30]
        if address:
            require(sheet in wb.sheetnames, f"Missing sheet: {sheet}")
            if cached:
                cache = load_workbook(path, data_only=True, read_only=True, keep_links=False)
                result["cache_note"] = "Stored results only; freshness unknown"
                cached_rows = list(cache[sheet].iter_rows(min_row=y1, max_row=y2, min_col=x1, max_col=x2, values_only=True))
            values = []
            for sh, addr, y, x in cells(sheet, address):
                c = wb[sh][addr]
                value = plan_value(c)
                if c.data_type == "f" and not isinstance(value, str):
                    value = {"type": type(value).__name__, **value.__dict__}
                row = {"cell": addr, "kind": c.data_type, "value": value, "number_format": c.number_format}
                if cached:
                    row["cached"] = cached_rows[y][x]
                values.append(row)
            result["cells"] = values
        return result
    finally:
        wb.close()
        if cache is not None:
            cache.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook")
    parser.add_argument("--sheet")
    parser.add_argument("--range", dest="address")
    parser.add_argument("--cached", action="store_true")
    parser.add_argument("--for-edit", action="store_true", help="Report replacement availability while continuing read-only inspection")
    args = parser.parse_args()
    print(json.dumps(inspect(args.workbook, args.sheet, args.address, args.cached, args.for_edit), ensure_ascii=True, default=str))
