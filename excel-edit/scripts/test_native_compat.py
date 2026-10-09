"""Development regressions: equivalent references, merged wrapping, failure evidence."""
from pathlib import Path
from tempfile import TemporaryDirectory
from copy import copy
import argparse
import json
import shutil
import zipfile
from lxml import etree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from edit_common import style, content, capture, load_plan, digest, EditError
from edit_native import canonical_matrix, execute
import edit_native
import native_preservation as preservation

def unit(root):
    def canon(s): return canonical_matrix([[s]],root)
    assert canon("='基础模型'!$H$35")==canon('=基础模型!$H$35')
    assert canon("='二期_行动与疲劳'!A1+1")==canon('=二期_行动与疲劳!A1+1')
    assert canon("='Bob''s 表'!A1")!=canon("='Bobs 表'!A1")
    assert canon('="\'表\'!A1"')!=canon('="表!A1"')
    assert canon("='表'!$A$1")!=canon('=表!A1')
    assert canon("='表'!A1")!=canon('=表2!A1')
    assert canon("=IF(A1,\"it's fine\",'表'!B1)")==canon('=IF(A1,"it\'s fine",表!B1)')
    w=Workbook();s=w.active;s.title='旧表';s.merge_cells('A1:B1');s['A1']='Label'
    metadata=preservation.capture_metadata(w)
    s.page_margins.left=2
    try:preservation.verify_metadata(w,metadata)
    except EditError:pass
    else:raise AssertionError('Changed page margin accepted')
    s.page_margins.left=metadata['旧表']['page_margins'].left
    s['A1'].alignment=Alignment(wrap_text=True,vertical='center')
    s['A2']='Unmerged';s['A2'].alignment=copy(s['A1'].alignment)
    before={(s.title,c.coordinate):(content(c),style(c)) for c in s._cells.values() if c.__class__.__name__!='MergedCell'}
    dp={'operations':[dict(sheet='旧表',range='A1',kind='value',before=[['Label']],values=[['Label']])]}
    captured=capture(w,dp)
    path=root/'repair.xlsx';w.save(path);w.close()
    with zipfile.ZipFile(path) as z: entries={i.filename:(i,z.read(i.filename)) for i in z.infolist()}
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    styles=ET.fromstring(entries['xl/styles.xml'][1])
    for a in styles.findall('m:cellXfs/m:xf/m:alignment',ns): a.attrib.pop('wrapText',None)
    info,_=entries['xl/styles.xml'];entries['xl/styles.xml']=(info,ET.tostring(styles))
    # Nonsequential worksheet path and namespace used only by an attribute value.
    info,data=entries.pop('xl/worksheets/sheet1.xml');info.filename='xl/worksheets/custom.xml'
    data=data.replace(b'<worksheet ',b'<worksheet xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac" mc:Ignorable="x14ac" ',1)
    entries[info.filename]=(info,data)
    info,data=entries['xl/_rels/workbook.xml.rels'];entries[info.filename]=(info,data.replace(b'sheet1.xml',b'custom.xml'))
    with zipfile.ZipFile(path,'w') as z:
        for info,data in entries.values():z.writestr(info,data)
    assert preservation.restore_merged_wrap(path,{'old':before},{'operations':[]})==1
    w=load_workbook(path)
    assert w.active['A1'].alignment.wrapText is True
    assert w.active['A2'].alignment.wrapText is None  # Never broaden repair to unmerged cells.
    assert style(w.active['A1'])==before[('旧表','A1')][1]
    w.close()
    differences=preservation.diagnose(path,dp,captured,{'old':before})
    assert any(d.get('location')=='旧表!A2' and d.get('property')=='alignment' for d in differences)
    with zipfile.ZipFile(path) as z:
        tree=ET.fromstring(z.read('xl/worksheets/custom.xml'))
        assert tree.nsmap['x14ac']=='http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac'
        assert z.read('xl/_rels/workbook.xml.rels')==entries['xl/_rels/workbook.xml.rels'][1]
    print('PASS unit: quote equivalence, literal/address differences, narrow wrap repair, relationships/namespaces',flush=True)

def native(root, fixture):
    path=root/'fixture.xlsx';shutil.copy2(fixture,path)
    w=load_workbook(path);name=w.sheetnames[1]
    assert 'B7:C7' in w[name].merged_cells
    before=preservation.capture_metadata(w);w.close()
    p=dict(workbook=str(path),new_sheets=[dict(name='兼容测试',range='A1:D3',row_height=26,
        formats=[dict(source_sheet=name,source_range='B7:C7',range='A1:B1')],merges=['A1:B1'])],
        operations=[dict(sheet='兼容测试',range='A1',kind='value',before=[[None]],values=[['较长的标签，用于验证写入之后保持声明的行高']]),
                    dict(sheet='兼容测试',range='D1',kind='formula',before=[[None]],values=[["='基础模型'!$H$35"]])],
        calculation='smart',acceptance=[dict(sheet='兼容测试',range='D1',values=[[.55]],tolerance=1e-9)])
    plan=root/'plan.json';plan.write_text(json.dumps(p),encoding='utf-8');p=load_plan(plan)
    original_verify=edit_native.verify_local
    def fail(*args,**kwargs): raise EditError('Intentional diagnostic failure')
    edit_native.verify_local=fail;h=digest(path)
    try:
        try:execute(p)
        except EditError as e:
            assert digest(path)==h and Path(e.context['staged_file']).exists()
            Path(e.context['staged_file']).unlink()
        else:raise AssertionError('Failure committed')
    finally:edit_native.verify_local=original_verify
    result=execute(p);print(json.dumps(result),flush=True)
    w=load_workbook(path);preservation.verify_metadata(w,before)
    assert abs(w['兼容测试'].row_dimensions[1].height-26)<0.2  # Native exact 26 verified by executor.
    w.close()
    print('PASS native: original task backup copy, row height 26, quoted reference, old metadata, failure retention',flush=True)

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--fixture');args=ap.parse_args()
    with TemporaryDirectory(prefix='excel-compat-') as d:
        root=Path(d);unit(root)
        if args.fixture:native(root,args.fixture)
