"""Create bounded sheets, copy selected formats, then use the ordinary native transaction."""
from copy import copy
import re
from edit_common import require, bounds, cells, style, content, effective_style_array, dimension_signature
from append_block import intersects


def validate(specs):
    require(isinstance(specs, list) and 1 <= len(specs) <= 5, 'new_sheets needs 1..5 definitions')
    names = set()
    for spec in specs:
        require(isinstance(spec, dict) and {'name', 'range'} <= set(spec) <=
                {'name', 'range', 'formats', 'merges', 'row_height', 'column_widths'}, 'Invalid new sheet fields')
        name = spec['name']
        require(isinstance(name, str) and 0 < len(name) <= 31 and not re.search(r"[\\/*?:\[\]]", name)
                and not name.startswith("'") and not name.endswith("'"), 'Invalid sheet name')
        require(name.casefold() not in names, 'Duplicate new sheet name')
        names.add(name.casefold())
        x1, y1, x2, y2 = bounds(spec['range'])
        require((x2-x1+1)*(y2-y1+1) <= 5000, 'New sheet region limited to 5,000 cells')
        def inside(address):
            a,b,c,d = bounds(address)
            require(x1 <= a <= c <= x2 and y1 <= b <= d <= y2, 'Region outside new sheet range')
            return a,b,c,d
        seen = set()
        for fmt in spec.get('formats', []):
            require(set(fmt) == {'source_sheet', 'source_range', 'range'}, 'Format copy needs source_sheet/source_range/range')
            a,b,c,d = bounds(fmt['source_range'])
            u,v,w,z = inside(fmt['range'])
            require((c-a,d-b) == (w-u,z-v), 'Format source/destination dimensions differ')
            keys = {addr for _,addr,_,_ in cells(name,fmt['range'])}
            require(not seen & keys, 'Overlapping format destinations')
            seen |= keys
        seen = set()
        for address in spec.get('merges', []):
            inside(address)
            keys = {addr for _,addr,_,_ in cells(name,address)}
            require(not seen & keys, 'Overlapping new merges')
            seen |= keys
        if 'row_height' in spec:
            require(type(spec['row_height']) in (int,float) and 0 < spec['row_height'] <= 409, 'Invalid row height')
        for col, width in spec.get('column_widths', {}).items():
            require(re.fullmatch('[A-Z]+',col) is not None, 'Column widths need single column letters')
            cx,_,_,_ = bounds(col+'1')
            require(x1 <= cx <= x2 and type(width) in (int,float) and 0 < width <= 255, f'Invalid column {col} width {width!r}: column must lie inside {spec["range"]}, width must be 0 < width <= 255')


def prepare(wb, p):
    specs = p.get('new_sheets')
    if not specs:
        return None
    validate(specs)
    require(sum(len(s._cells) for s in wb) <= 20000, 'Large sheet creation needs scoped handling')
    existing = set(wb.sheetnames)
    require(not {s['name'].casefold() for s in specs} & {n.casefold() for n in existing}, 'New sheet name already exists')
    old = {(s.title,c.coordinate):(content(c),style(c)) for s in wb for c in s._cells.values()
           if c.__class__.__name__ != 'MergedCell'}
    dims = {(s.title,kind,key):dimension_signature(d) for s in wb
            for kind in ('row_dimensions','column_dimensions') for key,d in getattr(s,kind).items()}
    merges = {s.title:str(s.merged_cells) for s in wb}
    result = dict(old=old, dims=dims, merges=merges, expected={})
    for spec in specs:
        ws = wb.create_sheet(spec['name'])
        for fmt in spec.get('formats', []):
            require(fmt['source_sheet'] in existing, 'Format source must be an existing sheet')
            source = wb[fmt['source_sheet']]
            sx,sy,ex,ey=bounds(fmt['source_range'])
            dx,dy,_,_=bounds(fmt['range'])
            from openpyxl.utils.cell import get_column_letter
            for m in source.merged_cells.ranges:
                if not intersects(fmt['source_range'],m): continue
                require(sx<=m.min_col<=m.max_col<=ex and sy<=m.min_row<=m.max_row<=ey,
                        'Format source cuts through a merged region')
                mapped=(f'{get_column_letter(dx+m.min_col-sx)}{dy+m.min_row-sy}:'
                        f'{get_column_letter(dx+m.max_col-sx)}{dy+m.max_row-sy}')
                require(mapped in spec.get('merges',[]), 'Declare the corresponding destination merge')
            require(not any(intersects(fmt['source_range'],r) for cf in source.conditional_formatting for r in cf.sqref.ranges),
                    'Conditional format source needs scoped handling')
            require(not any(intersects(fmt['source_range'],t.ref) for t in source.tables.values()), 'Table format source needs scoped handling')
            x,y,_,_ = bounds(fmt['source_range'])
            for _,addr,dy,dx in cells(ws.title,fmt['range']):
                ws[addr]._style = copy(effective_style_array(source.cell(y+dy,x+dx)))
        for address in spec.get('merges', []):
            ws.merge_cells(address)
        _,y1,_,y2 = bounds(spec['range'])
        if 'row_height' in spec:
            for r in range(y1,y2+1): ws.row_dimensions[r].height=spec['row_height']
        for col,width in spec.get('column_widths',{}).items():
            dimension=ws.column_dimensions[col]
            dimension.width=width
            dimension.min=dimension.max=bounds(col+'1')[0]
        result['expected'][ws.title] = {addr:style(ws[addr]) for _,addr,_,_ in cells(ws.title,spec['range'])
                                       if ws[addr].__class__.__name__ != 'MergedCell'}
    for op in p['operations']:
        if op['sheet'] not in existing:
            spec = next((s for s in specs if s['name'] == op['sheet']),None)
            require(spec is not None, 'Unknown operation sheet')
            a,b,c,d=bounds(op['range']); x,y,z,w=bounds(spec['range'])
            require(x<=a<=c<=z and y<=b<=d<=w,'Operation outside new sheet region')
    return result


def apply(target, specs):
    import pythoncom
    import win32com.client
    for spec in specs or []:
        collection=target.Worksheets
        last=collection.Item(collection.Count)
        # Explicit VARIANT omission is required here: this host's typed Add wrapper
        # drops the optional After argument. Raw dispatch preserves the sheet object.
        missing=win32com.client.VARIANT(pythoncom.VT_ERROR, -2147352572)
        raw=collection._oleobj_.Invoke(collection._oleobj_.GetIDsOfNames('Add'), 0,
                                      pythoncom.DISPATCH_METHOD, True, missing, last._oleobj_, 1)
        ws=win32com.client.Dispatch(raw)
        ws.Name=spec['name']
        for fmt in spec.get('formats',[]):
            target.Worksheets(fmt['source_sheet']).Range(fmt['source_range']).Copy()
            ws.Range(fmt['range']).PasteSpecial(Paste=-4122)
            target.Application.CutCopyMode=False
        for address in spec.get('merges',[]): ws.Range(address).Merge()
        if 'row_height' in spec: ws.Range(spec['range']).RowHeight=spec['row_height']
        for col,width in spec.get('column_widths',{}).items(): ws.Columns(col).ColumnWidth=width


def verify_native(target, specs):
    for spec in specs or []:
        ws=target.Worksheets(spec['name'])
        for col,width in spec.get('column_widths',{}).items():
            actual = ws.Columns(col).ColumnWidth
            require(abs(actual-width)<0.01,f'Saved Excel column width differs: {spec["name"]}/{col}, expected {width}, actual {actual}')
        if 'row_height' in spec:
            actual = ws.Range(spec['range']).RowHeight
            require(actual==spec['row_height'],f'Saved Excel row height differs: {spec["name"]}/{spec["range"]}, expected {spec["row_height"]}, actual {actual}')


def verify(wb,p,baseline,native=False):
    if baseline is None: return
    changed={(sh,addr) for o in p['operations'] for sh,addr,_,_ in cells(o['sheet'],o['range'])}
    append=p.get('append_block')
    appended={(sh,addr) for sh,addr,_,_ in cells(append['sheet'],append['range'])} if append else set()
    for (sh,addr),expected in baseline['old'].items():
        if (sh,addr) not in changed|appended:
            require((content(wb[sh][addr]),style(wb[sh][addr]))==expected, f'Old sheet changed: {sh}!{addr}')
    for (sh,kind,key),expected in baseline['dims'].items():
        if append and sh==append['sheet'] and kind=='row_dimensions':
            _,a,_,b=bounds(append['range'])
            if a<=key<=b: continue
        d=getattr(wb[sh],kind).get(key)
        require(d is not None and dimension_signature(d)==expected,f'Old dimension changed: {sh}/{kind}/{key}')
    for sh,expected in baseline['merges'].items():
        if not append or sh!=append['sheet']: require(str(wb[sh].merged_cells)==expected,'Old merges changed')
    for spec in p['new_sheets']:
        ws=wb[spec['name']]
        require({str(m) for m in ws.merged_cells.ranges}==set(spec.get('merges',[])), 'New sheet merges differ')
        for addr,expected in baseline['expected'][ws.title].items():
            if (ws.title,addr) not in changed:
                require(ws[addr].value is None and style(ws[addr])==expected,f'New region changed unexpectedly: {addr}')
        _,a,_,b=bounds(spec['range'])
        if 'row_height' in spec and not native:
            for r in range(a,b+1): require(ws.row_dimensions[r].height==spec['row_height'],'New row height differs')
