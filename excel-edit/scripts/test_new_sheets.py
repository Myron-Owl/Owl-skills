"""Temporary regression for new sheets, column styles and early lock reads."""
from pathlib import Path
from tempfile import TemporaryDirectory
import json
from copy import copy
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from edit_common import dimension_signature, load_plan, digest, EditError
from read_workbook import inspect
from edit_native import ExcelSession, execute
import new_sheets

def run(root):
    a,b=Workbook(),Workbook()
    b.active['A1'].font=Font(name='Arial',size=30)
    assert b.active['A1'].style_id>0
    for w in (a,b):
        w.active.column_dimensions['B'].font=Font(name='Arial',size=11)
        w.active.row_dimensions[2].font=Font(name='Arial',size=11)
    for kind,key in [('column_dimensions','B'),('row_dimensions',2)]:
        x,y=getattr(a.active,kind)[key],getattr(b.active,kind)[key]
        assert x.style_id != y.style_id
        assert dimension_signature(x)==dimension_signature(y)
    b.active.column_dimensions['B'].width=40
    assert dimension_signature(a.active.column_dimensions['B'])!=dimension_signature(b.active.column_dimensions['B'])
    b.active.row_dimensions[2].font=Font(name='Arial',size=30)
    assert dimension_signature(a.active.row_dimensions[2])!=dimension_signature(b.active.row_dimensions[2])
    a.close();b.close()
    path=root/'book.xlsx';w=Workbook();s=w.active;s.title='Base'
    s.append(['Label','Input','Result']);s.append(['Base',2,'=B2*2'])
    for row in s:
        for c in row:
            c.font=Font(name='Arial',size=11);c.fill=PatternFill('solid',fgColor='FFFFF2CC')
    s.row_dimensions[1].height=24;s.row_dimensions[2].height=24
    s.merge_cells('A4:C4');s['A4']='Title template'
    s['A4'].font=Font(name='Arial',size=20,bold=True);s.row_dimensions[4].height=30
    w.save(path);w.close()
    h=digest(path);owner=path.with_name('~$'+path.name);owner.write_text('test')
    try:
        result=inspect(path,for_edit=True)
        assert not result['edit_availability']['replaceable_now'] and digest(path)==h
    finally:owner.unlink()
    assert inspect(path,for_edit=True)['edit_availability']['replaceable_now']
    normalized=root/'native.xlsx'
    with ExcelSession() as session:
        book=session.open(path);book.SaveCopyAs(str(normalized));book.Close(False);book=None
    normalized.replace(path)
    def op(sheet,addr,before,value,kind='value'):
        return dict(sheet=sheet,range=addr,kind=kind,before=[[before]],values=[[value]])
    p=dict(workbook=str(path),new_sheets=[dict(name='Phase2',range='A1:C3',
        formats=[dict(source_sheet='Base',source_range='A1:C2',range='A1:C2'),
                 dict(source_sheet='Base',source_range='A4:C4',range='A3:C3')],
        merges=['A3:C3'],row_height=24,column_widths={'A':18})],
        append_block=dict(sheet='Base',range='A8:C8',row_height=24),
        operations=[op('Base','B2',2,3),op('Base','A8',None,'Added usage notes'),
                    op('Phase2','A1',None,'Phase2'),op('Phase2','B2',None,5),
                    op('Phase2','C2',None,'=Base!B2+B2','formula'),op('Phase2','A3',None,'Notes')],
        calculation='smart',acceptance=[dict(sheet='Base',range='C2',values=[[6]]),dict(sheet='Phase2',range='C2',values=[[8]])])
    plan=root/'plan.json';plan.write_text(json.dumps(p),encoding='utf-8');p=load_plan(plan)
    before=digest(path);bad=copy(p);bad['new_sheets']=[dict(name='base',range='A1:C3')]
    w=load_workbook(path)
    try:
        try:new_sheets.prepare(w,bad)
        except EditError:pass
        else:raise AssertionError('Duplicate name accepted')
    finally:w.close()
    assert digest(path)==before
    print(json.dumps(execute(p)),flush=True)
    w=load_workbook(path,data_only=True)
    assert w.sheetnames==['Base','Phase2'] and w['Phase2']['C2'].value==8
    assert w['Base']['C2'].value==6 and w['Base']['A8'].value=='Added usage notes'
    assert w['Phase2']['B2'].font.sz==11 and w['Phase2']['B2'].fill.fgColor.rgb=='FFFFF2CC'
    assert 'A3:C3' in w['Phase2'].merged_cells
    assert w['Phase2']['A3'].font.sz==20 and w['Phase2']['A3'].font.bold
    w.close()
    print('PASS: row/column indices, true format differences, early lock, new sheet + notes + formulas')

if __name__=='__main__':
    with TemporaryDirectory(prefix='excel-new-sheet-') as directory:run(Path(directory))
