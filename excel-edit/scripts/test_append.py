"""Synthetic small internal-formula workbook: append formats/merge/height and one input edit."""
from pathlib import Path
from tempfile import TemporaryDirectory
from copy import copy
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from edit_common import load_plan, digest, EditError, dimension_signature
from edit_native import execute
from append_block import prepare


def run(root):
    path = root / 'small.xlsx'
    w = Workbook()
    s = w.active
    s['A1'] = 'Input'
    s['B1'] = 'Value'
    s['C1'] = 'Result'
    s['A2'] = 'Base'
    s['B2'] = 2
    s['C2'] = '=B2*2'
    s.row_dimensions[1].height = 24
    s.row_dimensions[2].height = 24
    s.column_dimensions['A'].width = 18
    for row in s.iter_rows(min_row=1, max_row=2, max_col=3):
        for c in row:
            c.font = Font(name='Arial', size=11)
            c.fill = PatternFill('solid', fgColor='FFE6F0FA')
    w.save(path)
    # Establish an Excel-authored baseline, as in the user's existing workbook.
    from edit_native import ExcelSession
    normalized = root / 'normalized.xlsx'
    with ExcelSession() as session:
        book = session.open(path)
        book.SaveCopyAs(str(normalized))
        book.Close(SaveChanges=False)
        book = None
    normalized.replace(path)
    baseline = load_workbook(path)
    original_width = baseline.active.column_dimensions['A'].width
    baseline.close()
    # Equal-looking dimension styles with different table indices are equal.
    a = s.row_dimensions[3]
    a.font = Font(name='Arial', size=11)
    b = s.row_dimensions[4]
    b._style = copy(a._style)
    original = dimension_signature(a)[1]
    w._fonts.append(copy(w._fonts[b._style.fontId]))
    # IndexedList deduplicates normally; explicitly create a duplicate to emulate serialization.
    list.append(w._fonts, copy(w._fonts[b._style.fontId]))
    b._style.fontId = len(w._fonts)-1
    assert dimension_signature(b)[1] == original
    b.font = Font(name='Arial', size=30)
    assert dimension_signature(b)[1] != original
    w.close()
    p = dict(workbook=str(path), append_block=dict(sheet='Sheet', range='A8:C9',
             copy_formats_from='A1:C2', merges=['A8:B8'], row_height=24),
             operations=[dict(sheet='Sheet', range=addr, kind=kind, before=[[before]], values=[[after]])
                         for addr, before, after, kind in [('B2', 2, 3, 'value'), ('A8', None, 'New block', 'value'),
                             ('A9', None, 'Bonus', 'value'), ('B9', None, 5, 'value'), ('C9', None, '=B2+B9', 'formula')]],
             calculation='smart', acceptance=[dict(sheet='Sheet', range='C2', values=[[6]]),
                                               dict(sheet='Sheet', range='C9', values=[[8]])])
    plan_file = root / 'plan.json'
    plan_file.write_text(json.dumps(p), encoding='utf-8')
    validated = load_plan(plan_file)
    old_hash = digest(path)
    bad = copy(validated)
    bad['append_block'] = dict(sheet='Sheet', range='A2:C3')
    w = load_workbook(path)
    try:
        try:
            prepare(w, bad)
            raise AssertionError('Occupied append accepted')
        except EditError:
            assert digest(path) == old_hash
    finally:
        w.close()
    result = execute(validated)
    print(json.dumps(result), flush=True)
    w = load_workbook(path, data_only=True)
    s = w.active
    assert s['C2'].value == 6 and s['C9'].value == 8
    assert 'A8:B8' in s.merged_cells and s.row_dimensions[9].height == 24
    assert s['B9'].font.sz == 11 and s.row_dimensions[1].height == 24
    assert s['A1'].value == 'Input' and s.column_dimensions['A'].width == original_width
    w.close()
    print('PASS append: same session, values/formulas/results/formats/merge/heights/old cells/collision/style-index semantics')


if __name__ == '__main__':
    with TemporaryDirectory(prefix='excel-append-test-') as tmp:
        run(Path(tmp))
