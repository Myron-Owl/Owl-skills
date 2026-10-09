"""Shared bounded-edit contract. No Excel startup and no workbook writes on import."""
from __future__ import annotations

import argparse
from copy import copy
from datetime import date, datetime, time as time_value, timedelta
import hashlib
import json
import math
import os
from pathlib import Path
import re
import tempfile
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils.cell import range_boundaries, get_column_letter
from openpyxl.utils.datetime import to_excel


class EditError(RuntimeError):
    def __init__(self, message, **context):
        super().__init__(message)
        self.context = context



def require(ok, message):
    if not ok:
        raise EditError(message)


def digest(path):
    with open(path, "rb") as f:
        return hashlib.file_digest(f, "sha256").hexdigest()


def bounds(address):
    require(isinstance(address, str) and re.fullmatch(r"[A-Z]+[1-9][0-9]*(?::[A-Z]+[1-9][0-9]*)?", address),
            f"Use a bounded uppercase A1 address: {address!r}")
    x1, y1, x2, y2 = range_boundaries(address)
    require(1 <= x1 <= x2 <= 16384 and 1 <= y1 <= y2 <= 1048576, "Invalid Excel range")
    require((x2-x1+1)*(y2-y1+1) <= 100000, "Split edits larger than 100,000 cells into reviewed blocks")
    return x1, y1, x2, y2


def cells(sheet, address):
    x1, y1, x2, y2 = bounds(address)
    for y in range(y1, y2+1):
        for x in range(x1, x2+1):
            yield sheet, f"{get_column_letter(x)}{y}", y-y1, x-x1


def scalar(value):
    return value is None or type(value) in (str, bool, int) or (type(value) is float and math.isfinite(value))


def equal(a, b):
    # Excel represents all ordinary numbers as doubles, but bool and text remain distinct.
    if type(a) is bool or type(b) is bool:
        return type(a) is type(b) and a == b
    return a == b


def matrix(value, address):
    x1, y1, x2, y2 = bounds(address)
    require(isinstance(value, list) and len(value) == y2-y1+1, f"Wrong row count for {address}")
    require(all(isinstance(r, list) and len(r) == x2-x1+1 and all(scalar(v) for v in r) for r in value),
            f"Wrong columns or unsupported values for {address}")


def load_plan(filename):
    p = json.loads(Path(filename).read_text(encoding="utf-8-sig"))
    allowed = {"workbook", "operations", "preconditions", "calculation", "dependencies", "update_links", "expected_links", "acceptance", "timeout_seconds", "append_block", "new_sheets", "format_copies", "validations"}
    require(isinstance(p, dict) and not (set(p)-allowed), "Unknown plan fields")
    require(isinstance(p.get("workbook"), str) and Path(p["workbook"]).is_absolute(), "workbook must be an absolute path")
    p["workbook"] = str(Path(p["workbook"]).resolve(strict=True))
    require(isinstance(p.get("operations"), list) and (p["operations"] or p.get('format_copies') or p.get('validations')), "No operations")
    seen = set()
    total = 0
    for op in p["operations"]:
        require(isinstance(op, dict) and not(set(op)-{"sheet", "range", "kind", "values", "before", "color"}), "Unknown operation fields")
        require(isinstance(op.get("sheet"), str) and op["sheet"], "Missing sheet")
        bounds(op.get("range"))
        kind = op.get("kind")
        require(kind in ("value", "formula", "font_color", "fill_color"), "Unsupported operation kind")
        if kind in ("value", "formula"):
            require("before" in op and "values" in op and "color" not in op, "Content writes need before and values matrices")
            matrix(op["before"], op["range"])
            matrix(op["values"], op["range"])
            if kind == "formula":
                require(all(isinstance(v, str) and v.startswith("=") for row in op["values"] for v in row), "Formula writes must start with =")
        else:
            require(set(op) == {"sheet", "range", "kind", "color"}, "Color writes accept only sheet, range, kind, color")
            require(isinstance(op.get("color"), str) and re.fullmatch(r"#[0-9A-Fa-f]{6}", op["color"]), "Color must be #RRGGBB")
        for sh, addr, _, _ in cells(op["sheet"], op["range"]):
            key = (sh, addr, "content" if kind in ("value", "formula") else kind)
            require(key not in seen, f"Overlapping writes: {key}")
            seen.add(key)
            total += 1
    require(total <= 100000, "Plan exceeds 100,000 property writes; use a scoped custom batch")
    p.setdefault("preconditions", [])
    for check in p["preconditions"]:
        require(set(check) == {"sheet", "range", "values"}, "Precondition needs sheet/range/values")
        matrix(check["values"], check["range"])
    p.setdefault("calculation", "none")
    require(p["calculation"] in ("none", "smart", "full", "rebuild"), "Invalid calculation policy")
    p.setdefault("dependencies", [])
    require(isinstance(p["dependencies"], list), "dependencies must be a list")
    for dep in p["dependencies"]:
        require(isinstance(dep, str) and Path(dep).is_absolute() and Path(dep).is_file(), f"Missing dependency: {dep}")
    require(len({str(Path(d).resolve()).casefold() for d in p["dependencies"]}) == len(p["dependencies"]), "Duplicate dependencies")
    require(p["workbook"].casefold() not in {str(Path(d).resolve()).casefold() for d in p["dependencies"]}, "Target cannot also be a dependency")
    p.setdefault("update_links", False)
    require(type(p["update_links"]) is bool, "update_links must be Boolean")
    if "expected_links" in p:
        require(isinstance(p["expected_links"], list) and all(isinstance(v, str) and Path(v).is_absolute() for v in p["expected_links"]), "expected_links must contain absolute local workbook paths")
    p.setdefault("acceptance", [])
    for a in p["acceptance"]:
        require(set(a) <= {"sheet", "range", "values", "tolerance"} and {"sheet", "range", "values"} <= set(a), "Invalid acceptance fields")
        matrix(a["values"], a["range"])
        require(type(a.get("tolerance", 0)) in (int, float) and math.isfinite(a.get("tolerance", 0)) and a.get("tolerance", 0) >= 0, "Invalid tolerance")
    if p["calculation"] != "none":
        require(bool(p["acceptance"]), "Calculation needs explicit expected result cells")
    require(not p["update_links"] or (p["calculation"] != "none" and p["dependencies"]), "Link update needs dependencies and calculated acceptance")
    require(not p["dependencies"] or p["update_links"], "Opening sources can refresh linked values; set update_links=true explicitly")
    p.setdefault("timeout_seconds", 180)
    require(type(p["timeout_seconds"]) is int and 10 <= p["timeout_seconds"] <= 3600, "timeout_seconds must be 10..3600")
    if "append_block" in p:
        from append_block import validate
        validate(p["append_block"])
    if "new_sheets" in p:
        from new_sheets import validate
        validate(p["new_sheets"])
    return p


def package_features(path):
    require(Path(path).suffix.lower() in (".xlsx", ".xlsm"), "Fixed editors support .xlsx/.xlsm; use a scoped native workflow for other formats")
    require(zipfile.is_zipfile(path), "Encrypted or invalid workbook; do not bypass encryption")
    flags = set()
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        for n in names:
            low = n.lower()
            for marker, label in (("externallinks/", "external_links"), ("vbaproject", "macros"), ("connections.xml", "connections"),
                                  ("querytables/", "connections"), ("model/", "data_model"), ("pivot", "pivots"),
                                  ("drawings/", "drawings"), ("charts/", "charts"), ("activex/", "activex"),
                                  ("slicer", "extensions"), ("signature", "signatures"), ("macrosheets/", "xlm_macros"),
                                  ("customxml/", "custom_xml")):
                if marker in low:
                    flags.add(label)
            if n.startswith("xl/worksheets/") and n.endswith(".xml") or n == "xl/workbook.xml":
                # One bounded package pass selects the route; no hash inventory or artifact report.
                with z.open(n) as f:
                    for _, e in ET.iterparse(f, events=("end",)):
                        local = e.tag.rsplit("}", 1)[-1]
                        if local == "f": flags.add("formulas")
                        if local in ("sheetProtection", "workbookProtection") and any(v not in ("0", "false", "False", "") for v in e.attrib.values()): flags.add("protection")
                        if local == "extLst": flags.add("extensions")
                        e.clear()
    return flags


def check_unlocked(path):
    require(not Path(path).with_name("~$"+Path(path).name).exists(), "Workbook has an Excel owner file; close it before replacing it")
    if os.name == "nt":
        import ctypes
        from ctypes import wintypes
        k = ctypes.WinDLL("kernel32", use_last_error=True)
        k.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        k.CreateFileW.restype = wintypes.HANDLE
        k.CloseHandle.argtypes = [wintypes.HANDLE]
        h = k.CreateFileW(str(path), 0x80000000 | 0x40000000, 0, None, 3, 0x80, None)
        require(h != ctypes.c_void_p(-1).value, f"Workbook is locked or not writable (Windows error {ctypes.get_last_error()})")
        k.CloseHandle(h)


def stage_path(path):
    fd, name = tempfile.mkstemp(prefix=".excel-edit-", suffix=Path(path).suffix, dir=Path(path).parent)
    os.close(fd)
    return Path(name)


def commit(temp, original, old_hash):
    check_unlocked(original)
    require(digest(original) == old_hash, "Source changed during editing; original kept")
    os.replace(temp, original)


def read_book(path):
    return load_workbook(path, data_only=False, keep_links=True, rich_text=True)


def content(cell):
    v = cell.value
    if cell.data_type == "f" and not isinstance(v, str):
        return (cell.data_type, type(v).__name__, dict(v.__dict__))
    return (cell.data_type, v)


def plan_value(cell):
    value = cell.value
    if isinstance(value, CellRichText):
        return str(value)  # Matching only; never flatten the stored rich text.
    if isinstance(value, (datetime, date, time_value, timedelta)):
        return to_excel(value, epoch=cell.parent.parent.epoch)
    return value


def formula_map(wb):
    # Used only for native edits: one in-memory pass, no COM per-cell formula audit.
    return {(ws.title, c.coordinate): content(c) for ws in wb for c in ws._cells.values() if c.data_type == "f"}


def effective_style_array(cell):
    if cell._style is not None:
        return cell._style
    ws = cell.parent
    row = ws.row_dimensions.get(cell.row)
    if row is not None and row.has_style:
        return row._style
    for col in ws.column_dimensions.values():
        if col.has_style and (col.min or 0) <= cell.column <= (col.max or 0):
            return col._style
    # Excel's Normal style need not be an all-zero StyleArray (e.g. localized templates).
    return ws.parent._cell_styles[0]


def style(cell):
    c = copy(cell)
    c._style = effective_style_array(cell)
    return {k: copy(getattr(c, k)) for k in ("font", "fill", "border", "alignment", "protection")} | {"number_format": c.number_format}


def dimension_signature(dimension):
    """Compare resolved properties, never workbook-local style indices."""
    attrs = dict(dimension)
    attrs.pop('s', None)
    attrs.pop('style', None)
    return attrs, {k: copy(getattr(dimension, k)) for k in
                   ('font', 'fill', 'border', 'alignment', 'protection', 'number_format')}


def capture(wb, p):
    targets = {}
    rows, cols, merges = {}, {}, {}
    changes = {}
    for op in p["operations"]:
        require(op["sheet"] in wb.sheetnames, f"Missing sheet: {op['sheet']}")
        ws = wb[op["sheet"]]
        require(not ws.protection.sheet, f"Protected target sheet: {ws.title}")
        for sh, addr, _, _ in cells(op["sheet"], op["range"]):
            c = ws[addr]
            require(c.__class__.__name__ != "MergedCell", f"Merged non-anchor cell: {sh}!{addr}")
            if op["kind"].endswith("color"):
                require(not any(c.coordinate in cf.sqref for cf in ws.conditional_formatting),
                        f"Conditional formatting controls {sh}!{addr}; inspect its rule instead of forcing a static color")
            key = (sh, addr)
            targets.setdefault(key, (content(c), style(c)))
            changes.setdefault(key, set()).add(op["kind"])
            rd = ws.row_dimensions.get(c.row)
            rows[(sh, c.row)] = None if rd is None else (rd.height, rd.hidden, rd.outlineLevel, rd.collapsed)
            # Dimensions covering grouped columns are included; absent entries stay absent.
            if sh not in cols:
                cols[sh] = {k: (d.min, d.max, d.width, d.hidden, d.outlineLevel, d.collapsed) for k, d in ws.column_dimensions.items()}
                merges[sh] = str(ws.merged_cells)
    return dict(targets=targets, changes=changes, rows=rows, cols=cols, merges=merges, sheets=wb.sheetnames)


def table_header_issues(wb):
    issues = []
    for ws in wb:
        for table in ws.tables.values():
            if table.headerRowCount == 0:
                continue
            x1, y1, x2, _ = range_boundaries(table.ref)
            if len(table.tableColumns) != x2-x1+1:
                issues.append(dict(sheet=ws.title, table=table.name, range=table.ref,
                                   issue="column_count_mismatch"))
                continue
            for x, column in enumerate(table.tableColumns, x1):
                cell = ws.cell(y1, x)
                visible = plan_value(cell)
                if not isinstance(visible, str) or visible != column.name:
                    issues.append(dict(sheet=ws.title, table=table.name, cell=cell.coordinate,
                                       stored=column.name, visible=visible))
    return issues


def check_tables(wb, p=None):
    issues = table_header_issues(wb)
    require(not issues, "Table header metadata mismatch; inspect a copy before repair: " +
            json.dumps(issues[:8], ensure_ascii=True, default=str))
    if p is None:
        return
    for op in p["operations"]:
        if op["kind"] not in ("value", "formula") or op["sheet"] not in wb.sheetnames:
            continue
        for table in wb[op["sheet"]].tables.values():
            if table.headerRowCount == 0:
                continue
            tx1, ty1, tx2, _ = range_boundaries(table.ref)
            x1, y1, x2, y2 = bounds(op["range"])
            require(not (y1 <= ty1 <= y2 and x1 <= tx2 and tx1 <= x2),
                    f"Table header edit needs scoped column/reference handling: {table.name}")


def check_before(wb, p):
    check_tables(wb, p)
    checks = list(p["preconditions"]) + [dict(sheet=o["sheet"], range=o["range"], values=o["before"]) for o in p["operations"] if "before" in o]
    for ck in checks:
        require(ck["sheet"] in wb.sheetnames, f"Missing sheet: {ck['sheet']}")
        for sh, addr, y, x in cells(ck["sheet"], ck["range"]):
            actual = plan_value(wb[sh][addr])
            require(equal(actual, ck["values"][y][x]), f"Before-value conflict at {sh}!{addr}: expected {ck['values'][y][x]!r}, found {actual!r}")


def verify_local(wb, p, baseline, native=False):
    check_tables(wb)
    current = capture(wb, p)
    for key in ("rows", "cols", "merges", "sheets"):
        actual, expected = current[key], baseline[key]
        if key == 'rows' and native and p.get('new_sheets'):
            # Native RowHeight is checked after reopening. OOXML can record 26.1 for 26.
            names = {s['name'] for s in p['new_sheets'] if 'row_height' in s}
            actual = {k: ((None, *v[1:]) if k[0] in names and v else v) for k,v in actual.items()}
            expected = {k: ((None, *v[1:]) if k[0] in names and v else v) for k,v in expected.items()}
        if key == 'cols' and native and p.get('new_sheets'):
            # New widths are specified in Excel ColumnWidth units, while package widths
            # include padding. Verify those through Excel after reopen, not raw equality.
            names = {s['name'] for s in p['new_sheets']}
            actual = {k:v for k,v in actual.items() if k not in names}
            expected = {k:v for k,v in expected.items() if k not in names}
        detail = (f': actual {actual!r}, expected {expected!r}' if not isinstance(actual, dict) else
                  ': ' + repr([(k, expected.get(k), actual.get(k)) for k in dict.fromkeys([*expected, *actual]) if expected.get(k) != actual.get(k)][:8]))
        require(actual == expected, f"Unrequested layout change: {key}{detail}")
    for key, (old_content, old_style) in baseline["targets"].items():
        new_content, new_style = current["targets"][key]
        allowed = baseline["changes"][key]
        if not allowed.intersection({"value", "formula"}):
            require(new_content == old_content, f"Color-only edit changed content at {key}")
        for component, before in old_style.items():
            after = new_style[component]
            if component == "font" and "font_color" in allowed:
                before, after = copy(before), copy(after)
                before.color = after.color = None
            if component == "fill" and "fill_color" in allowed:
                continue
            require(before == after, f"Unrequested {component} change at {key}")
    for op in p["operations"]:
        for sh, addr, y, x in cells(op["sheet"], op["range"]):
            c = wb[sh][addr]
            kind = op["kind"]
            if kind in ("value", "formula"):
                value = op["values"][y][x]
                if not (native and kind == "formula"):
                    require(equal(plan_value(c), value), f"Saved content mismatch at {sh}!{addr}: {c.value!r}")
                require((c.data_type == "f") == (kind == "formula"), f"Wrong formula/value kind at {sh}!{addr}")
                if type(value) is str and kind == "value":
                    require(c.data_type in ("s", "inlineStr"), f"Text converted at {sh}!{addr}")
            else:
                col = c.font.color if kind == "font_color" else c.fill.fgColor
                require(col is not None and col.type == "rgb" and col.rgb[-6:].upper() == op["color"][1:].upper(), f"Saved color mismatch at {sh}!{addr}")
                if kind == "fill_color": require(c.fill.patternType == "solid", f"Not a solid fill at {sh}!{addr}")


def main_cli(execute, description):
    parser = argparse.ArgumentParser(description=description)
    parser.add_argument("plan", help="UTF-8 JSON plan; see references/edit-plan.md")
    args = parser.parse_args()
    p = None
    try:
        p = load_plan(args.plan)
        result = execute(p)
        print(json.dumps(result, ensure_ascii=True))
    except Exception as e:
        print(json.dumps({"status": "FAILED", "error": str(e), "stage": "preflight",
                          "workbook": p["workbook"] if p else None,
                          "original_written": False, **getattr(e, "context", {})}, ensure_ascii=True))
        raise SystemExit(1)
