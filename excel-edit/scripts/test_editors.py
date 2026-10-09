"""One-time regression fixtures. Run after editor changes, never per user workbook.

Usage: excel-python.cmd test_editors.py --output ABSOLUTE_EMPTY_DIRECTORY [--native]
All Excel files are synthetic and stay in the specified directory.
"""
import argparse
from copy import copy
import json
from pathlib import Path
import sys
import time

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils.datetime import CALENDAR_WINDOWS_1900, CALENDAR_MAC_1904, to_excel

from edit_common import load_plan, digest, EditError, content, style, package_features
import edit_simple


def plan(root, filename, workbook, operations, **kw):
    p = dict(workbook=str(workbook), operations=operations, **kw)
    file = root/filename
    file.write_text(json.dumps(p, ensure_ascii=False), encoding="utf-8")
    return load_plan(file)


def op(sheet, cell, before, after, kind="value"):
    return dict(sheet=sheet, range=cell, kind=kind, before=[[before]], values=[[after]])


def reject_unchanged(execute, p):
    before = digest(p["workbook"])
    try:
        execute(p)
    except EditError as error:
        assert digest(p["workbook"]) == before
        retained = error.context.get('staged_file')
        if retained:
            assert Path(retained).exists()
            Path(retained).unlink()
    else:
        raise AssertionError("Unsafe edit was not rejected")


def simple_tests(root):
    file = root/"简单测试.xlsx"
    w = Workbook()
    s = w.active
    s.title = "数据"
    s.append(["标题", "说明", "数量"])
    s.append(["0012", "旧文字", 10])
    s.append(["0013", "不要动", 20])
    s["A1"].font = Font(name="Arial", size=28, bold=True)
    s["B2"].font = Font(name="Arial", size=11, italic=True)
    s["B2"].alignment = Alignment(wrap_text=True)
    s.row_dimensions[1].height = 45
    s.row_dimensions[2].height = 24
    s.column_dimensions["B"].width = 19
    s.merge_cells("E1:F1")
    s["E1"] = "合并标题"
    s["D3"] = "=只是文字"
    s["D3"].data_type = "s"
    w.save(file)
    w.close()
    original = load_workbook(file)
    untouched = {(ws.title, c.coordinate): (content(c), style(c)) for ws in original for c in ws._cells.values() if c.coordinate not in {"A2", "B2", "C2"}}
    original.close()
    p = plan(root, "simple.json", file, [op("数据", "A2", "0012", "0014"), op("数据", "B2", "旧文字", "=这不是公式"), op("数据", "C2", 10, 12),
        dict(sheet="数据", range="B2", kind="font_color", color="#FF0000"), dict(sheet="数据", range="B2", kind="fill_color", color="#FFFF00")])
    print(edit_simple.execute(p), flush=True)
    w = load_workbook(file)
    s = w["数据"]
    assert s["A2"].value == "0014" and s["B2"].data_type == "s"
    assert s["B2"].font.sz == 11 and s["B2"].font.italic and s.row_dimensions[2].height == 24
    assert s.row_dimensions[1].height == 45 and s.column_dimensions["B"].width == 19
    for (sh, addr), expected in untouched.items(): assert (content(w[sh][addr]), style(w[sh][addr])) == expected
    w.close()
    reject_unchanged(edit_simple.execute, p)  # Old-value conflict.
    reject_unchanged(edit_simple.execute, plan(root, "merge.json", file, [op("数据", "F1", None, "bad")]))
    w = load_workbook(file)
    w["数据"]["D2"] = "=C2*2"
    w.save(file)
    w.close()
    reject_unchanged(edit_simple.execute, plan(root, "formula_reject.json", file, [op("数据", "C2", 12, 14)]))
    cf = root/"条件格式.xlsx"
    w = Workbook()
    w.active["A1"] = 2
    w.active.conditional_formatting.add("A1", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FF00FF00")))
    w.save(cf)
    w.close()
    reject_unchanged(edit_simple.execute, plan(root, "cf.json", cf, [dict(sheet="Sheet", range="A1", kind="fill_color", color="#FFFF00")]))
    for label, epoch in [("1900", CALENDAR_WINDOWS_1900), ("1904", CALENDAR_MAC_1904)]:
        date_file = root/("日期"+label+".xlsx")
        w = Workbook()
        w.epoch = epoch
        w.active["A1"].number_format = "yyyy-mm-dd"
        w.save(date_file)
        w.close()
        edit_simple.execute(plan(root, "date-new-"+label+".json", date_file, [op("Sheet", "A1", None, 46000)]))
        edit_simple.execute(plan(root, "date-edit-"+label+".json", date_file, [op("Sheet", "A1", 46000, 46001)]))
        w = load_workbook(date_file)
        assert to_excel(w.active["A1"].value, epoch=w.epoch) == 46001
        assert w.active["A1"].number_format == "yyyy-mm-dd"
        w.close()
    print("PASS simple: content/types/colors/unrelated styles/layout/conflicts/merged cell/formula routing/conditional format", flush=True)


def native_tests(root):
    from edit_native import ExcelSession, execute
    source = root/"数值源.xlsx"
    target = root/"引用表.xlsx"
    with ExcelSession() as session:
        a = session.app
        src = a.Workbooks.Add()
        src.Worksheets(1).Name = "参数"
        src.Worksheets(1).Range("A1").Value2 = 7
        src.SaveAs(str(source), FileFormat=51)
        dst = a.Workbooks.Add()
        s = dst.Worksheets(1)
        s.Name = "数据"
        s.Range("A1").Value2 = "大标题"
        s.Range("A1").Font.Size = 28
        s.Rows(1).RowHeight = 45
        s.Range("B2").Value2 = 3
        s.Range("B2").Font.Size = 11
        s.Rows(2).RowHeight = 24
        s.Columns(2).ColumnWidth = 19
        s.Range("C2").Formula2 = "=B2*2"
        s.Range("D2").Formula2 = "='[数值源.xlsx]参数'!$A$1+B2"
        s.Range("E2").Value2 = "'0012"
        s.Range("G2").Formula2 = "=SEQUENCE(2)"
        s.Range("J1:K1").Merge()
        s.Range("J1").Value2 = "合并标题"
        s.Range("A4").Formula2 = "=NA()"  # Existing error is not a reason to repair the entire workbook.
        s.Range("H6").NumberFormat = "yyyy-mm-dd"
        a.Calculate()
        a.Calculation = -4105
        dst.SaveAs(str(target), FileFormat=51)
        src.Close(False)
        dst.Close(False)
        src = dst = s = a = None
    source_hash = digest(source)
    old = load_workbook(target)
    original_formula = old["数据"]["D2"].value
    original_calc_mode = old.calculation.calcMode
    old.close()
    p = plan(root, "native_color.json", target, [dict(sheet="数据", range="B2", kind="font_color", color="#FF0000"), dict(sheet="数据", range="B2", kind="fill_color", color="#FFFF00")])
    print(execute(p), flush=True)
    saved = load_workbook(target)
    assert saved.calculation.calcMode == original_calc_mode
    assert saved["数据"]["D2"].value == original_formula
    assert saved["数据"]["B2"].font.sz == 11 and saved["数据"].row_dimensions[2].height == 24
    saved.close()
    p = plan(root, "native_calc.json", target, [op("数据", "B2", 3, 5), op("数据", "C2", "=B2*2", "=B2*3", "formula"), op("数据", "E2", "0012", "0015"), op("数据", "F2", None, "=文字"),
        dict(sheet="数据", range="A6:B7", kind="value", before=[[None,None],[None,None]], values=[[10,True],["0018","=literal"]]),
        dict(sheet="数据", range="C6:C7", kind="formula", before=[[None],[None]], values=[["=A6+1"],["=LEN(A7)"]]), op("数据", "H6", None, 46000)],
        calculation="smart", dependencies=[str(source)], update_links=True,
        acceptance=[dict(sheet="数据", range="C2:D2", values=[[15,12]]), dict(sheet="数据", range="C6:C7", values=[[11],[4]])])
    print(execute(p), flush=True)
    w = load_workbook(target, data_only=True)
    assert w["数据"]["C2"].value == 15 and w["数据"]["D2"].value == 12
    assert w["数据"]["E2"].value == "0015" and w["数据"]["F2"].value == "=文字"
    w.close()
    assert digest(source) == source_hash
    reject_unchanged(execute, plan(root, "wrong_result.json", target, [op("数据", "B2", 5, 6)], calculation="smart", dependencies=[str(source)], update_links=True, acceptance=[dict(sheet="数据", range="C2", values=[[999]])]))
    reject_unchanged(execute, plan(root, "spill_target.json", target, [op("数据", "G3", 2, 99)]))
    w = load_workbook(target)
    before_ref = w["数据"]["D2"].value
    w.close()
    ext = "='"+str(root).replace("'", "''")+"\\[数值源.xlsx]参数'!$A$1+B2+1"
    print(execute(plan(root, "external_formula.json", target, [op("数据", "D2", before_ref, ext, "formula")],
                      calculation="smart", dependencies=[str(source)], update_links=True,
                      acceptance=[dict(sheet="数据", range="D2", values=[[13]])])), flush=True)
    assert digest(source) == source_hash
    print("PASS native: source-closed reopen/formula calculation/external formula edits/links/unchanged source/dynamic formula preservation/layout/text types/failed acceptance keeps source", flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True)
    parser.add_argument("--native", action="store_true")
    args = parser.parse_args()
    root = Path(args.output)
    if not root.is_absolute() or (root.exists() and any(root.iterdir())):
        raise SystemExit("Use an absolute empty test directory")
    root.mkdir(parents=True, exist_ok=True)
    start = time.monotonic()
    simple_tests(root)
    if args.native: native_tests(root)
    assert not list(root.glob(".excel-edit-*")), "Staging files leaked"
    print(f"ALL REQUESTED TESTS PASSED in {time.monotonic()-start:.1f}s", flush=True)
